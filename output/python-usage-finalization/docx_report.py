from pathlib import Path
from docx import Document

output = Path("report.docx")
document = Document()
document.add_heading("Quarterly report", level=0)
document.add_paragraph("Café results — created with python-docx 1.2.0.")
table = document.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Item"
table.rows[0].cells[1].text = "Count"
row = table.add_row().cells
row[0].text, row[1].text = "Reports", "3"
document.save(output)
reopened = Document(output)
reopened.tables[0].cell(1, 1).text = "4"
reopened.save(output)
assert Document(output).tables[0].cell(1, 1).text == "4"
with output.open("rb") as handle:
    assert handle.read(2) == b"PK"
print(output.name, "reopened and edited")
