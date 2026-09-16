from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import xlsxwriter

output = Path("report.xlsx")
book = Workbook()
sheet = book.active
sheet.title = "Summary"
sheet.append(["Count", "Value"])
sheet.append([1, 7])
sheet.append([2, 3])
sheet["B4"] = "=SUM(B2:B3)"
sheet["A1"].font = Font(bold=True)
book.save(output)
book.close()
edited = load_workbook(output)
edited["Summary"]["B2"] = 9
edited.save(output)
edited.close()
check = load_workbook(output)
assert check["Summary"]["B2"].value == 9
assert check["Summary"]["B4"].value == "=SUM(B2:B3)"
check.close()
cached = load_workbook(output, data_only=True)
assert cached["Summary"]["B4"].value is None
cached.close()

with xlsxwriter.Workbook("writer.xlsx") as writer:
    sheet = writer.add_worksheet("Summary")
    sheet.write_row("A1", [7, 3])
    # The caller supplies 10; XlsxWriter does not calculate this formula.
    sheet.write_formula("C1", "=SUM(A1:B1)", None, 10)
cached = load_workbook("writer.xlsx", data_only=True)
assert cached["Summary"]["C1"].value == 10
cached.close()
print("report.xlsx edited; writer.xlsx cached result verified")
