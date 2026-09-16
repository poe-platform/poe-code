"""Ordinary Python script; browser qualification currently uses runpy.run_path.

PROBE_ROOT chooses canonical safe-bash storage or the MEMFS control. PROBE_FONT
is the staged, licensed TTF fixture. This does not qualify shell python FILE
dispatch. No desktop-Python execution is evidence.
"""

import datetime
import errno
import hashlib
import importlib.metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import traceback
import zipfile


ROOT = Path(os.environ["PROBE_ROOT"])
ROOT.mkdir(parents=True, exist_ok=True)
WORK = ROOT / "documents with spaces ünicode 日本語"
WORK.mkdir(exist_ok=True)
DEFAULT_TEMP_DIRECTORY = tempfile.gettempdir()
tempfile.tempdir = str(WORK)
RESULT = {"root": str(ROOT), "workflows": {}, "artifacts": {}, "versions": {}, "default_temp_directory": DEFAULT_TEMP_DIRECTORY, "qualified_temp_directory": str(WORK)}
for distribution in (
    "python-docx", "openpyxl", "XlsxWriter", "pypdf", "fpdf2", "lxml",
    "Pillow", "fonttools", "defusedxml", "et-xmlfile", "typing-extensions",
):
    try:
        RESULT["versions"][distribution] = importlib.metadata.version(distribution)
    except importlib.metadata.PackageNotFoundError:
        RESULT["versions"][distribution] = None


def record_artifact(path):
    payload = path.read_bytes()
    assert payload, str(path)
    RESULT["artifacts"][str(path.relative_to(ROOT))] = {
        "bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest()
    }


def check_zip(path, required_parts):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        for part in required_parts:
            assert any(name.startswith(part) for name in names), (part, names)
        return names


def filesystem_probe():
    target = WORK / "seek file café.bin"
    with target.open("w+b") as stream:
        stream.write(b"abcdef")
        stream.flush()
        assert stream.seek(2) == 2
        assert stream.tell() == 2
        stream.write(b"XY")
        stream.flush()
        stream.seek(-2, os.SEEK_END)
        assert stream.read() == b"ef"
        stream.seek(0)
        assert stream.read() == b"abXYef"
        stream.truncate(4)
    assert target.read_bytes() == b"abXY"
    memory = io.BytesIO(b"abcdef")
    memory.seek(-3, os.SEEK_END)
    assert memory.read() == b"def"
    large = WORK / "large boundary ü.bin"
    payload = bytes(range(256)) * 385
    large.write_bytes(payload)
    assert large.read_bytes() == payload
    with large.open("r+b") as stream:
        stream.seek(65530)
        stream.write(b"cross boundary")
        stream.flush()
        stream.seek(65530)
        assert stream.read(14) == b"cross boundary"
    assert large.stat().st_size == len(payload)
    record_artifact(large)
    with tempfile.TemporaryDirectory(prefix="temporary ü ") as directory:
        assert Path(directory).parent == WORK
        with tempfile.NamedTemporaryFile(dir=directory, suffix=" λ.bin") as stream:
            stream.write(b"temporary bytes")
            stream.flush()
            stream.seek(0)
            assert stream.read() == b"temporary bytes"
            named = Path(stream.name)
            assert named.read_bytes() == b"temporary bytes"
        assert not named.exists()
    assert not Path(directory).exists()
    record_artifact(target)
    return {"assertions": ["Unicode/space paths", "seek/tell/end-relative seek", "truncate", "BytesIO", "98560-byte read/write and 64KiB boundary seek", "named temp reopen/unlink", "temporary directory cleanup"]}


def image_probe():
    from PIL import Image, features
    target = WORK / "image fixture ü.png"
    Image.new("RGB", (24, 16), (24, 96, 160)).save(target)
    with Image.open(target) as image:
        assert image.size == (24, 16)
        assert image.getpixel((0, 0)) == (24, 96, 160)
    record_artifact(target)
    return {"size": [24, 16], "native_features": {name: features.version(name) for name in features.get_supported()}}


def nofollow_exclusive_probe():
    directory = WORK / "nofollow exclusive ü"
    directory.mkdir(exist_ok=True)
    regular = directory / "existing regular"
    regular.write_bytes(b"preserve existing bytes")
    target_directory = directory / "target directory"
    target_directory.mkdir(exist_ok=True)
    (target_directory / "child").write_bytes(b"preserve child bytes")
    absent = directory / "absent target"
    links = {
        "dangling link": absent,
        "regular link": regular,
        "directory link": target_directory,
        "self loop": directory / "self loop",
    }
    for name, target in links.items():
        os.symlink(str(target), str(directory / name))
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW
    outcomes = {}
    for name in ("fresh file", "existing regular", "target directory", *links):
        try:
            descriptor = os.open(directory / name, flags, 0o600)
        except OSError as error:
            outcomes[name] = error.errno
        else:
            try:
                os.write(descriptor, b"created exclusive bytes")
            finally:
                os.close(descriptor)
            outcomes[name] = "created"
    # Characterize the deliberately distinct nonexclusive capability. MEMFS
    # supports regular-file O_NOFOLLOW; the bridge still refuses it as ENOTSUP.
    try:
        descriptor = os.open(regular, os.O_RDONLY | os.O_NOFOLLOW)
    except OSError as error:
        outcomes["nonexclusive_regular"] = error.errno
    else:
        os.close(descriptor)
        outcomes["nonexclusive_regular"] = "opened"
    expected_nonexclusive = "opened" if os.environ.get("PROBE_PROFILE") == "memfs" else errno.ENOTSUP
    expected = {"fresh file": "created", "existing regular": errno.EEXIST, "target directory": errno.EEXIST, **{name: errno.EEXIST for name in links}, "nonexclusive_regular": expected_nonexclusive}
    assert outcomes == expected, {"actual": outcomes, "expected": expected}
    assert (directory / "fresh file").read_bytes() == b"created exclusive bytes"
    assert regular.read_bytes() == b"preserve existing bytes"
    assert (target_directory / "child").read_bytes() == b"preserve child bytes"
    assert not absent.exists()
    for name, target in links.items():
        assert os.readlink(directory / name) == str(target)
    return {"assertions": ["exclusive nofollow fresh create", "existing regular/directory and dangling/regular/directory/self-loop symlinks EEXIST", "targets and symlinks preserved", "nonexclusive nofollow capability unchanged"], "outcomes": outcomes}


def docx_probe():
    from docx import Document
    from docx.shared import Inches
    from lxml import etree
    target = WORK / "heading table café.docx"
    document = Document()
    document.add_heading("Heading café", level=1)
    document.add_paragraph("Original paragraph 日本語")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Label"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Before"
    table.cell(1, 1).text = "42"
    document.add_picture(str(WORK / "image fixture ü.png"), width=Inches(0.5))
    document.save(target)
    reopened = Document(target)
    assert reopened.paragraphs[0].text == "Heading café"
    assert reopened.paragraphs[0].style.name == "Heading 1"
    assert reopened.paragraphs[1].text == "Original paragraph 日本語"
    assert reopened.tables[0].cell(1, 1).text == "42"
    assert len(reopened.inline_shapes) == 1
    reopened.paragraphs[0].text = "Edited heading café"
    reopened.paragraphs[1].text = "Edited paragraph 日本語"
    reopened.tables[0].cell(1, 0).text = "After"
    reopened.inline_shapes[0].width = Inches(0.75)
    reopened.add_paragraph("Appended paragraph")
    reopened.save(target)
    memory = io.BytesIO()
    Document(target).save(memory)
    memory.seek(0)
    final = Document(memory)
    assert final.paragraphs[0].text == "Edited heading café"
    assert final.paragraphs[0].style.name == "Heading 1"
    assert final.paragraphs[1].text == "Edited paragraph 日本語"
    assert final.paragraphs[-1].text == "Appended paragraph"
    assert final.tables[0].cell(1, 0).text == "After"
    assert len(final.inline_shapes) == 1
    assert final.inline_shapes[0].width == Inches(0.75)
    parts = check_zip(target, ["word/document.xml", "word/media/image", "word/_rels/document.xml.rels"])
    with zipfile.ZipFile(target) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        assert "Edited heading café" in xml.xpath("//w:t/text()", namespaces=namespaces)
        assert len(xml.xpath("//w:tbl", namespaces=namespaces)) == 1
        assert len(xml.xpath("//a:blip", namespaces=namespaces)) == 1
    record_artifact(target)
    return {"assertions": ["heading/paragraph/table/image create reopen edit", "BytesIO reopen", "ZIP CRC and image/relationship parts"], "parts": len(parts), "libxml_version": list(etree.LIBXML_VERSION), "libxslt_version": list(etree.LIBXSLT_VERSION)}


def openpyxl_probe():
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font, PatternFill
    target = WORK / "edited dates ü.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data café"
    sheet.append(["Value", "Date", "Formula"])
    sheet.append([7, datetime.datetime(2026, 1, 2), "=A2*2"])
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="123456")
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet.add_image(Image(str(WORK / "image fixture ü.png")), "E1")
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "E5")
    workbook.save(target)
    reopened = load_workbook(target)
    sheet = reopened["Data café"]
    assert sheet["A2"].value == 7
    assert sheet["B2"].value == datetime.datetime(2026, 1, 2)
    assert sheet["C2"].value == "=A2*2"
    assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb == "00123456"
    assert sheet["B2"].number_format == "yyyy-mm-dd"
    sheet["A2"] = 9
    sheet["B2"] = datetime.datetime(2026, 2, 3)
    sheet["C2"] = "=A2*3"
    sheet["A2"].font = Font(italic=True)
    reopened.save(target)
    memory = io.BytesIO(target.read_bytes())
    final = load_workbook(memory)["Data café"]
    assert final["A2"].value == 9 and final["A2"].font.italic
    assert final["B2"].value == datetime.datetime(2026, 2, 3)
    assert final["C2"].value == "=A2*3"
    assert final["A1"].font.bold and final["A1"].fill.fgColor.rgb == "00123456"
    assert final["B2"].number_format == "yyyy-mm-dd"
    assert len(final._images) == 1 and len(final._charts) == 1
    assert load_workbook(target, data_only=True)["Data café"]["C2"].value is None
    parts = check_zip(target, ["xl/media/image", "xl/charts/chart", "xl/drawings/drawing", "xl/worksheets/sheet1.xml"])
    record_artifact(target)
    return {"assertions": ["values/dates/formulas/styles create reopen edit", "BytesIO read", "image/chart survive edit", "formula stored; no recalculation/cached result"], "parts": len(parts)}


def xlsxwriter_probe():
    import xlsxwriter
    from openpyxl import load_workbook
    target = WORK / "formatted multiple sheets 日本語.xlsx"
    with xlsxwriter.Workbook(str(target), {"tmpdir": str(WORK)}) as workbook:
        data = workbook.add_worksheet("Data café")
        summary = workbook.add_worksheet("Summary")
        heading = workbook.add_format({"bold": True, "bg_color": "#123456", "font_color": "#FFFFFF"})
        date = workbook.add_format({"num_format": "yyyy-mm-dd"})
        data.write_row(0, 0, ["Value", "Date", "Formula"], heading)
        data.write_number(1, 0, 11)
        data.write_datetime(1, 1, datetime.datetime(2026, 3, 4), date)
        data.write_formula(1, 2, "=A2*2", None, 22)
        data.set_column("A:C", 20)
        data.freeze_panes(1, 0)
        data.insert_image("E1", str(WORK / "image fixture ü.png"))
        summary.write("A1", "Summary café", heading)
        chart = workbook.add_chart({"type": "column"})
        chart.add_series({"values": "='Data café'!$A$2:$A$2"})
        summary.insert_chart("A3", chart)
    reopened = load_workbook(target)
    assert reopened.sheetnames == ["Data café", "Summary"]
    sheet = reopened["Data café"]
    assert sheet["A2"].value == 11
    assert sheet["B2"].value == datetime.datetime(2026, 3, 4)
    assert sheet["C2"].value == "=A2*2"
    assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb == "FF123456"
    assert sheet["B2"].number_format == "yyyy-mm-dd"
    assert sheet.freeze_panes == "A2"
    assert load_workbook(target, data_only=True)["Data café"]["C2"].value == 22
    assert load_workbook(io.BytesIO(target.read_bytes()))["Summary"]["A1"].value == "Summary café"
    parts = check_zip(target, ["xl/media/image", "xl/charts/chart", "xl/worksheets/sheet2.xml", "xl/drawings/_rels/"])
    memory = io.BytesIO()
    with xlsxwriter.Workbook(memory, {"in_memory": True}) as workbook:
        workbook.add_worksheet().write("A1", "BytesIO café")
    memory.seek(0)
    assert load_workbook(memory).active["A1"].value == "BytesIO café"
    record_artifact(target)
    return {"assertions": ["formatted multiple sheets", "dates/formulas/image/chart parts", "openpyxl interoperability", "explicit cached formula 22; no recalculation", "BytesIO write/read", "disk-backed temporary packaging"], "parts": len(parts)}


def pdf_probe():
    import fpdf
    from fpdf import FPDF
    from pypdf import PdfReader, PdfWriter
    assert fpdf.__version__ == importlib.metadata.version("fpdf2")
    try:
        legacy = importlib.metadata.version("fpdf")
    except importlib.metadata.PackageNotFoundError:
        legacy = None
    assert legacy is None, "Legacy fpdf distribution conflicts with fpdf2"
    target = WORK / "multi page café.pdf"
    pdf = FPDF()
    pdf.set_title("Document probe café")
    pdf.set_author("safe-bash runtime qualification")
    pdf.add_font("ProbeFont", fname=os.environ["PROBE_FONT"])
    for page in range(1, 4):
        pdf.add_page()
        pdf.set_font("ProbeFont", size=14)
        pdf.cell(text=f"Page {page} café embedded font")
        pdf.image(str(WORK / "image fixture ü.png"), x=20, y=30, w=24)
    pdf.output(str(target))
    reader = PdfReader(target)
    assert len(reader.pages) == 3
    assert reader.metadata.title == "Document probe café"
    assert reader.metadata.author == "safe-bash runtime qualification"
    for index, page in enumerate(reader.pages, 1):
        assert f"Page {index} café embedded font" in page.extract_text()
        resources = page["/Resources"]
        assert any(obj.get_object().get("/Subtype") == "/Image" for obj in resources["/XObject"].values())
        fonts = [obj.get_object() for obj in resources["/Font"].values()]
        descendants = [font["/DescendantFonts"][0].get_object() for font in fonts if "/DescendantFonts" in font]
        assert descendants
        for font in descendants:
            descriptor = font["/FontDescriptor"].get_object()
            assert len(descriptor["/FontFile2"].get_object().get_data()) > 100
        for image in resources["/XObject"].values():
            image = image.get_object()
            if image.get("/Subtype") == "/Image":
                assert image["/Width"] == 24 and image["/Height"] == 16
                assert image.get_data()
    split = WORK / "split page ü.pdf"
    writer = PdfWriter()
    writer.add_page(reader.pages[1])
    writer.write(split)
    assert len(PdfReader(split).pages) == 1
    assert "Page 2 café" in PdfReader(split).pages[0].extract_text()
    merged = WORK / "merged round trip ü.pdf"
    writer = PdfWriter()
    writer.append(target)
    writer.append(split)
    writer.add_metadata({"/Title": "Merged café"})
    writer.write(merged)
    final = PdfReader(io.BytesIO(merged.read_bytes()))
    assert len(final.pages) == 4 and final.metadata.title == "Merged café"
    assert [f"Page {n} café" in page.extract_text() for n, page in zip([1, 2, 3, 2], final.pages)] == [True] * 4
    for page in final.pages:
        resources = page["/Resources"]
        assert any(image.get_object().get("/Subtype") == "/Image" and image.get_object().get_data() for image in resources["/XObject"].values())
        fonts = [font.get_object() for font in resources["/Font"].values()]
        assert any(font["/DescendantFonts"][0].get_object()["/FontDescriptor"].get_object()["/FontFile2"].get_object().get_data() for font in fonts if "/DescendantFonts" in font)
    memory = io.BytesIO()
    writer.write(memory)
    memory.seek(0)
    assert len(PdfReader(memory).pages) == 4
    assert len(PdfReader(io.BytesIO(bytes(pdf.output()))).pages) == 3
    for path in (target, split, merged):
        record_artifact(path)
    return {"assertions": ["fpdf2 distribution identity; legacy fpdf absent", "3 pages; embedded TTF/image resources", "Unicode text/metadata extraction", "split/merge/roundtrip", "BytesIO seek write/read"]}


def pymupdf_probe():
    import pymupdf
    target = WORK / "multi page café.pdf"
    with pymupdf.open(str(target)) as document:
        assert len(document) == 3
        assert "Page 1 café" in document[0].get_text()
        rich = document[0].get_text("dict")
        assert any(block["type"] == 0 for block in rich["blocks"])
        assert any(block["type"] == 1 for block in rich["blocks"])
        pixmap = document[0].get_pixmap(matrix=pymupdf.Matrix(0.5, 0.5))
        assert pixmap.width > 100 and pixmap.height > 100
        rendered = WORK / "rendered page ü.png"
        pixmap.save(str(rendered))
        record_artifact(rendered)
    RESULT["versions"]["PyMuPDF"] = importlib.metadata.version("PyMuPDF")
    return {"assertions": ["page rendering PNG", "structured text and image blocks"], "pixels": [pixmap.width, pixmap.height]}


def stream_modes_probe():
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfReader, PdfWriter
    import xlsxwriter

    # Real retained binary handles, not only path strings or BytesIO.
    target = WORK / "stream edit café.docx"
    with target.open("w+b") as stream:
        document = Document()
        document.add_paragraph("Stream café 日本語")
        document.save(stream)
        stream.seek(0)
        assert Document(stream).paragraphs[0].text == "Stream café 日本語"
    with tempfile.SpooledTemporaryFile(max_size=64, mode="w+b", dir=WORK) as stream:
        document.save(stream)
        assert stream._rolled
        stream.seek(0)
        assert Document(stream).paragraphs[0].text == "Stream café 日本語"
    written = WORK / "write only ü.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Stream ü")
    for index in range(150):
        sheet.append([index, f"row {index} 日本語", f"=A{index + 1}+1"])
    workbook.save(written)
    with written.open("rb") as stream:
        reopened = load_workbook(stream, read_only=True)
        try:
            rows = list(reopened.active.values)
            assert rows == [(index, f"row {index} 日本語", f"=A{index + 1}+1") for index in range(150)]
        finally:
            reopened.close()
    constant = WORK / "constant memory ü.xlsx"
    with xlsxwriter.Workbook(str(constant), {"constant_memory": True, "tmpdir": str(WORK)}) as workbook:
        sheet = workbook.add_worksheet()
        for index in range(150):
            sheet.write_row(index, 0, [index, f"row {index} 日本語"])
    reopened = load_workbook(constant, read_only=True)
    try:
        assert list(reopened.active.values) == [(index, f"row {index} 日本語") for index in range(150)]
    finally:
        reopened.close()
    pdf_path = WORK / "stream roundtrip ü.pdf"
    with (WORK / "multi page café.pdf").open("rb") as source, pdf_path.open("w+b") as destination:
        reader = PdfReader(source)
        writer = PdfWriter()
        writer.add_page(reader.pages[2])
        writer.write(destination)
        destination.seek(0)
        assert "Page 3 café" in PdfReader(destination).pages[0].extract_text()
    for path in (target, written, constant, pdf_path):
        record_artifact(path)
    return {"assertions": ["DOCX retained w+b handle", "DOCX spooled temporary rollover", "openpyxl write-only and read-only 150 rows", "XlsxWriter constant-memory 150 rows", "pypdf retained input/output handles"]}


def invalid_inputs_probe():
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    from PIL import Image, UnidentifiedImageError

    outcomes = {}
    for name, loader, expected in (
        ("docx", Document, zipfile.BadZipFile),
        ("xlsx", load_workbook, zipfile.BadZipFile),
        ("pdf", PdfReader, PdfReadError),
        ("image", Image.open, UnidentifiedImageError),
    ):
        for payload_name, payload in (("empty", b""), ("garbage", b"not a document\x00\xff")):
            target = WORK / f"invalid {name} {payload_name}.bin"
            target.write_bytes(payload)
            with target.open("rb") as stream:
                try:
                    loader(stream)
                except expected as error:
                    outcomes[f"{name}_{payload_name}"] = type(error).__name__
                else:
                    raise AssertionError(f"{name} accepted {payload_name}")
            assert target.read_bytes() == payload
            target.unlink()
    # A rejected parse must not poison subsequent library or file operations.
    assert Document(WORK / "heading table café.docx").paragraphs[0].text == "Edited heading café"
    assert load_workbook(WORK / "edited dates ü.xlsx").active["A2"].value == 9
    assert len(PdfReader(WORK / "multi page café.pdf").pages) == 3
    return {"assertions": ["empty/corrupt DOCX XLSX PDF PNG rejected with library errors", "input bytes preserved", "valid reopen after failure"], "outcomes": outcomes}


def timestamp_probe():
    import shutil
    source = WORK / "timestamps source ü.bin"
    target = WORK / "timestamps copied ü.bin"
    source.write_bytes(b"timestamp contents")
    os.utime(source, (1600000000, 1700000000))
    observed = source.stat()
    assert observed.st_atime == 1600000000 and observed.st_mtime == 1700000000
    shutil.copy2(source, target)
    assert target.stat().st_mtime == 1700000000
    assert target.read_bytes() == b"timestamp contents"
    source.unlink()
    target.unlink()
    return {"assertions": ["distinct atime/mtime set and read back", "shutil.copy2 content and mtime preservation"]}


def runtime_resources_probe():
    import importlib.resources
    for package in ("importlib", "encodings", "docx"):
        source = importlib.resources.files(package).joinpath("__init__.py").read_text(encoding="utf-8")
        assert len(source) > 100, package
    return {"assertions": ["resources from already-loaded standard-library ZIP packages", "python-docx package resources"]}


def root_namespace_probe():
    previous_temp = tempfile.tempdir
    try:
        tempfile.tempdir = None
        assert tempfile.gettempdir() == "/tmp"
        with tempfile.NamedTemporaryFile(prefix="default ü ", delete=False) as stream:
            stream.write(b"canonical default temporary bytes")
            temporary = Path(stream.name)
        assert temporary.parent == Path("/tmp")
        assert temporary.read_bytes() == b"canonical default temporary bytes"
    finally:
        tempfile.tempdir = previous_temp
    link = WORK / "absolute link outside work"
    link.symlink_to(temporary)
    assert link.read_bytes() == temporary.read_bytes()
    link.unlink()
    Path("/lib").mkdir(exist_ok=True)
    user_file = Path("/lib/user owned ü.txt")
    user_file.write_bytes(b"canonical user lib bytes")
    assert user_file.read_bytes() == b"canonical user lib bytes"
    # The parent must independently compare these out-of-/work bytes against
    # its canonical filesystem; worker-local visibility alone is insufficient.
    return {"assertions": ["default temporary files use canonical /tmp", "absolute symlink outside /work", "canonical user /lib stays visible"], "artifacts": [{"path": str(path), "bytes": list(path.read_bytes())} for path in (temporary, user_file)]}


for name, probe in (
    ("filesystem", filesystem_probe), ("pillow_image", image_probe),
    ("nofollow_exclusive", nofollow_exclusive_probe),
    ("python_docx", docx_probe), ("openpyxl", openpyxl_probe),
    ("xlsxwriter", xlsxwriter_probe), ("fpdf2_pypdf", pdf_probe),
    ("stream_modes", stream_modes_probe), ("invalid_inputs", invalid_inputs_probe),
    ("timestamps", timestamp_probe),
    ("runtime_resources", runtime_resources_probe),
):
    try:
        RESULT["workflows"][name] = {"status": "passed", **probe()}
    except Exception as error:
        RESULT["workflows"][name] = {"status": "failed", "error": str(error), "traceback": traceback.format_exc()}

if os.environ.get("PROBE_PROFILE") == "root-bridge":
    try:
        RESULT["workflows"]["root_namespace"] = {"status": "passed", **root_namespace_probe()}
    except Exception as error:
        RESULT["workflows"]["root_namespace"] = {"status": "failed", "error": str(error), "traceback": traceback.format_exc()}

RESULT["priority_passed"] = all(item["status"] == "passed" for item in RESULT["workflows"].values())
if importlib.util.find_spec("pymupdf") is None:
    RESULT["pymupdf"] = {"status": "unresolved", "reason": "PyMuPDF not installed; selected Pyodide index must be evaluated by the browser harness"}
else:
    try:
        RESULT["pymupdf"] = {"status": "passed", **pymupdf_probe()}
    except Exception as error:
        RESULT["pymupdf"] = {"status": "failed", "error": str(error), "traceback": traceback.format_exc()}
RESULT["reportlab"] = {
    "status": "not_needed" if RESULT["workflows"]["fpdf2_pypdf"]["status"] == "passed" else "generation_failure_requires_evaluation",
    "reason": "Only evaluate ReportLab when a required PDF generation feature is missing; extraction/rendering are separate concerns",
}
(ROOT / "report.json").write_text(json.dumps(RESULT, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(RESULT, ensure_ascii=False))
if not RESULT["priority_passed"]:
    raise SystemExit(1)
