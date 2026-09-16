"""Ordinary guest program: all artifacts live on the canonical filesystem."""
import io
import os
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.drawing.image import Image as WorkbookImage
import xlsxwriter
from PIL import Image
from fpdf import FPDF

assert Path('host-input.txt').read_text() == 'canonical host input\n'
with tempfile.NamedTemporaryFile(dir='/tmp', mode='w+b') as temporary:
    temporary.write(b'0123456789')
    temporary.seek(4)
    temporary.write(b'AB')
    temporary.seek(-3, os.SEEK_END)
    assert temporary.read() == b'789'
    temporary.truncate(8)
    temporary.seek(0)
    Path('random-access.bin').write_bytes(temporary.read())
assert not Path(temporary.name).exists()
# A solid diagnostic pixel fixture is not artwork.
Image.new('RGB', (80, 40), color=(30, 100, 170)).save('fixture.png')

doc = Document()
doc.add_heading('Public Python document verification', 0)
doc.add_paragraph('Created by python-docx through the canonical filesystem.')
doc.add_heading('Verified table', level=1)
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Item'
table.rows[0].cells[1].text = 'Value'
cells = table.add_row().cells
cells[0].text, cells[1].text = 'Canonical input', '42'
doc.add_picture('fixture.png', width=Inches(1))
doc.save('report.docx')

book = Workbook()
sheet = book.active
sheet.title = 'Results'
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_setup.orientation = 'landscape'
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
sheet.print_area = 'A1:L20'
sheet.append(['Item', 'Value'])
sheet.append(['First', 2])
sheet.append(['Second', 3])
sheet['B4'] = '=SUM(B2:B3)'
sheet['A1'].font = Font(bold=True, color='FFFFFF')
sheet['A1'].fill = PatternFill('solid', fgColor='1E64AA')
sheet.column_dimensions['A'].width = 24
chart = BarChart()
chart.title = 'Canonical values'
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
sheet.add_chart(chart, 'D2')
sheet.add_image(WorkbookImage('fixture.png'), 'A7')
book.save('openpyxl.xlsx')

with xlsxwriter.Workbook('writer.xlsx') as writer:
    sheet = writer.add_worksheet('Results')
    sheet.set_landscape()
    sheet.fit_to_pages(1, 1)
    sheet.print_area('A1:L20')
    header = writer.add_format({'bold': True, 'bg_color': '#1E64AA', 'font_color': '#FFFFFF'})
    sheet.write_row('A1', ['Item', 'Value'], header)
    sheet.write_row('A2', ['First', 2])
    sheet.write_row('A3', ['Second', 3])
    sheet.write_formula('B4', '=SUM(B2:B3)', None, 5)
    sheet.set_column('A:A', 24)
    chart = writer.add_chart({'type': 'column'})
    chart.add_series({'values': '=Results!$B$2:$B$3'})
    sheet.insert_chart('D2', chart)
    sheet.insert_image('A7', 'fixture.png')

pdf = FPDF()
pdf.add_page()
pdf.set_font('Helvetica', size=16)
pdf.cell(text='Public Python PDF verification')
pdf.image('fixture.png', x=10, y=30, w=40)
pdf.add_font('FixtureMono', fname='font.ttf')
pdf.add_page()
pdf.set_font('FixtureMono', size=12)
pdf.cell(text='Second page before edit')
pdf.output('report.pdf')
Path('guest-output.txt').write_text('created by guest\n')
print('created')
