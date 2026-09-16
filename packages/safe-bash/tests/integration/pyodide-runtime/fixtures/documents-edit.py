from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

assert Path('host-edit.txt').read_text() == 'host reopened artifacts\n'
doc = Document('report.docx')
doc.tables[0].rows[1].cells[1].text = '73'
doc.add_paragraph('Reopened and edited in a fresh interpreter.')
doc.save('report.docx')
for filename in ['openpyxl.xlsx', 'writer.xlsx']:
    book = load_workbook(filename)
    book['Results']['B2'] = 7
    book.save(filename)
reader = PdfReader('report.pdf')
writer = PdfWriter()
writer.add_page(reader.pages[0])
writer.add_metadata({'/Title': 'Reopened canonical PDF'})
writer.write('edited.pdf')
Path('guest-output.txt').write_text('edited by guest\n')
print('edited')
