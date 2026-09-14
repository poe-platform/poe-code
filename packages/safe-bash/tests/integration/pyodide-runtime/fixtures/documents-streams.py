"""Reopen documents through binary streams and names with spaces/non-ASCII."""
import tempfile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

folder = Path('résultats with spaces')
folder.mkdir()
with open('report.docx', 'rb') as source:
    document = Document(source)
with open(folder / 'edited résumé.docx', 'w+b') as target:
    document.save(target)
    target.seek(0)
    assert Document(target).tables[0].rows[1].cells[1].text == '73'
for filename in ['openpyxl.xlsx', 'writer.xlsx']:
    # A spool forced onto canonical /tmp exercises seekable temporary storage.
    with tempfile.SpooledTemporaryFile(max_size=1, mode='w+b', dir='/tmp') as spool:
        with open(filename, 'rb') as source:
            workbook = load_workbook(source)
        workbook.save(spool)
        spool.seek(0)
        reopened = load_workbook(spool)
        assert reopened['Results']['B2'].value == 7
        assert reopened['Results']['B4'].value == '=SUM(B2:B3)'
        assert len(reopened['Results']._charts) == 1
        assert len(reopened['Results']._images) == 1
        spool.seek(0)
        (folder / filename).write_bytes(spool.read())
with open('edited.pdf', 'rb') as source:
    reader = PdfReader(source)
    writer = PdfWriter()
    writer.append(reader)
    with open(folder / 'edited résumé.pdf', 'w+b') as target:
        writer.write(target)
        target.seek(0)
        reopened = PdfReader(target)
        assert len(reopened.pages) == 1
        assert len(reopened.pages[0].images) == 1
        assert 'Public Python PDF verification' in reopened.pages[0].extract_text()
assert not list(Path('/tmp').iterdir())
print('streams unicode paths spooled temporary files')
