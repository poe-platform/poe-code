import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

assert Path('random-access.bin').read_bytes() == b'0123AB67'
doc = Document('report.docx')
assert doc.paragraphs[0].text == 'Public Python document verification'
assert any(p.style.name == 'Heading 1' for p in doc.paragraphs)
assert doc.tables[0].rows[1].cells[1].text == '73'
assert len(doc.inline_shapes) == 1
assert doc.paragraphs[-1].text == 'Reopened and edited in a fresh interpreter.'
# Independent OOXML parsing complements reopening with the producing library.
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile('report.docx') as archive:
    xml = ET.fromstring(archive.read('word/document.xml'))
    assert len(xml.findall('.//w:tbl', ns)) == 1
    assert '73' in [node.text for node in xml.findall('.//w:t', ns)]
    assert len([name for name in archive.namelist() if name.startswith('word/media/')]) == 1
for filename in ['openpyxl.xlsx', 'writer.xlsx']:
    book = load_workbook(filename)
    sheet = book['Results']
    assert sheet['B2'].value == 7
    assert sheet['B4'].value == '=SUM(B2:B3)'
    assert sheet['A1'].font.bold
    assert sheet['A1'].fill.fgColor.rgb == '001E64AA' or sheet['A1'].fill.fgColor.rgb == 'FF1E64AA'
    assert len(sheet._charts) == 1
    assert len(sheet._images) == 1
    # openpyxl stores formulas, but does not calculate them on save.
    assert load_workbook(filename, data_only=True)['Results']['B4'].value is None
    with ZipFile(filename) as archive:
        names = archive.namelist()
        assert any(name.startswith('xl/charts/chart') for name in names)
        assert any(name.startswith('xl/media/') for name in names)
        xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        assert xml.find('.//{*}f').text == 'SUM(B2:B3)'
original, edited = PdfReader('report.pdf'), PdfReader('edited.pdf')
assert len(original.pages) == 2
assert len(edited.pages) == 1
assert 'Second page before edit' in original.pages[1].extract_text()
assert 'Public Python PDF verification' in edited.pages[0].extract_text()
assert len(edited.pages[0].images) == 1
fonts = edited.pages[0]['/Resources']['/Font'].get_object()
assert any(font.get_object()['/BaseFont'] == '/Helvetica' for font in fonts.values())
embedded = original.pages[1]['/Resources']['/Font'].get_object()
assert any('/FontFile2' in child.get_object()['/FontDescriptor'].get_object()
           for font in embedded.values()
           for child in font.get_object().get('/DescendantFonts', []))
assert edited.metadata.title == 'Reopened canonical PDF'
print(json.dumps({'docx': 'heading table image edited', 'xlsx': 'values styles formulas charts images edited', 'pdf': 'text pages image Helvetica edited', 'io': 'temporary seek truncate'}))
